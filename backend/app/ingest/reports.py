"""Immutable report ingestion and bounded, provenance-preserving parsers."""
from __future__ import annotations
import hashlib, io, json, re, uuid, zipfile
from datetime import date
from pathlib import Path, PurePosixPath
from sqlalchemy import and_, or_, select
from sqlalchemy.orm import Session
from ..db.models import FileRecord, Fragment, Project, Report
from ..settings import get_settings

MAX_BYTES = 10 * 1024 * 1024
MAX_PDF_PAGES = 20
MAX_XLSX_ROWS = 5000
MAX_EXTRACTED_CHARS = 2_000_000
MAX_ZIP_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_ZIP_MEMBERS = 2048
MAX_XLSX_COLUMNS = 200
MAX_XLSX_SHEETS = 30
SAFE_NAME = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._ -]{0,254}$")

def _normalise(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def validate_document_archive(raw: bytes) -> None:
    """Bound OOXML containers before a library opens individual members.

    No archive member is extracted to disk. Reject ambiguous or escaping member
    names anyway, so downstream readers see one canonical package inventory.
    """
    if len(raw) > MAX_BYTES:
        raise ValueError("report exceeds 10 MB limit")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ZIP_MEMBERS:
                raise ValueError("document contains too many archive members")
            if sum(item.file_size for item in members) > MAX_ZIP_EXPANDED_BYTES:
                raise ValueError("document expands beyond 50 MB limit")
            seen = set()
            for item in members:
                name = item.filename
                path = PurePosixPath(name)
                if ("\\" in name or path.is_absolute() or ".." in path.parts
                        or (path.parts and ":" in path.parts[0])
                        or name in seen or item.flag_bits & 1
                        or (item.external_attr >> 16) & 0o170000 == 0o120000):
                    raise ValueError("unsafe or ambiguous document archive member")
                seen.add(name)
    except zipfile.BadZipFile as exc:
        raise ValueError("invalid document archive") from exc

def _filename(name: str | None) -> str:
    name = Path(name or "pasted-report.txt").name
    if not SAFE_NAME.fullmatch(name) or name in {".", ".."}:
        raise ValueError("unsafe filename")
    return name


def _store_original(storage_key: str, raw: bytes) -> None:
    """Write immutable bytes outside the static tree using an atomic rename."""
    root = Path(get_settings().upload_dir).resolve()
    target = root / storage_key
    if root not in target.parents:
        raise ValueError("invalid storage key")
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_suffix(target.suffix + ".pending")
    try:
        with temporary.open("xb") as handle:
            handle.write(raw)
            handle.flush()
        temporary.replace(target)
    except Exception:
        # A failed write must not leave a resumable-looking upload behind.
        # Keep cleanup best-effort so the original exception remains visible.
        for path in (temporary, target):
            try:
                path.unlink(missing_ok=True)
            except OSError:
                pass
        raise


def _remove_original(storage_key: str) -> None:
    """Remove only the artifact created by the current intake attempt."""
    root = Path(get_settings().upload_dir).resolve()
    target = (root / storage_key).resolve()
    if root not in target.parents:
        return
    for path in (target.with_suffix(target.suffix + ".pending"), target):
        try:
            path.unlink(missing_ok=True)
        except OSError:
            # Compensation must not hide the database error that caused it.
            pass

def parse_bytes(raw: bytes, filename: str, mime: str) -> tuple[list[tuple[str, str]], list[str], str]:
    from .report_parser import run_parser
    result = run_parser('parse', raw, filename=filename, mime=mime)
    return [tuple(part) for part in result['parts']], result['warnings'], result['kind']


def _parse_bytes_local(raw: bytes, filename: str, mime: str) -> tuple[list[tuple[str, str]], list[str], str]:
    """Pure parser implementation, called only inside the isolated worker."""
    if len(raw) > MAX_BYTES:
        raise ValueError("report exceeds 10 MB limit")
    suffix = Path(filename).suffix.lower()
    warnings: list[str] = []
    extracted_chars = 0
    def add_part(parts: list[tuple[str, str]], locator: str, value: str) -> None:
        nonlocal extracted_chars
        extracted_chars += len(value)
        if extracted_chars > MAX_EXTRACTED_CHARS:
            raise ValueError("extracted report text exceeds limit")
        parts.append((locator, value))
    if suffix in {".docx", ".xlsx"}:
        validate_document_archive(raw)
    if suffix in ("", ".txt") or (suffix not in {".pdf", ".xlsx", ".docx"} and mime.startswith("text/")):
        try: text = raw.decode("utf-8-sig")
        except UnicodeDecodeError as exc: raise ValueError("text file must be UTF-8") from exc
        if not text.strip(): raise ValueError("report is empty")
        parts=[]
        for i, paragraph in enumerate(re.split(r"\n\s*\n", text), 1):
            if paragraph.strip(): add_part(parts, f"paragraph:{i}", paragraph)
        if not parts: raise ValueError("report is empty")
        return parts, warnings, "text"
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc: raise ValueError("DOCX support is unavailable") from exc
        try: doc = Document(io.BytesIO(raw))
        except Exception as exc: raise ValueError("invalid DOCX document") from exc
        parts=[]
        from docx.text.paragraph import Paragraph
        paragraph_number = table_number = 0
        for block in doc.iter_inner_content():
            if isinstance(block, Paragraph):
                paragraph_number += 1
                if block.text.strip():
                    add_part(parts, f"paragraph:{paragraph_number}", block.text)
            else:
                table_number += 1
                for ri, row in enumerate(block.rows, 1):
                    text = " | ".join(c.text for c in row.cells)
                    if text.strip():
                        add_part(parts, f"table:{table_number}/row:{ri}", text)
        if not parts: raise ValueError("DOCX contains no readable text")
        return parts, warnings, "docx"
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc: raise ValueError("PDF support is unavailable") from exc
        try:
            reader=PdfReader(io.BytesIO(raw))
            if reader.is_encrypted:
                raise ValueError("encrypted PDF requires an unlocked source")
        except Exception as exc:
            raise ValueError("invalid or encrypted PDF document") from exc
        if len(reader.pages)>MAX_PDF_PAGES: raise ValueError("PDF exceeds 20 page limit")
        parts=[]
        for i,page in enumerate(reader.pages,1):
            text=(page.extract_text() or "").strip()
            if len(re.sub(r"\s+", "", text)) < 40:
                warnings.append(f"NEEDS_TRANSCRIPTION:page:{i}")
            add_part(parts, f"page:{i}", text)
        if not parts: raise ValueError("PDF contains no extractable text")
        return parts,warnings,"pdf"
    if suffix in (".xlsx",):
        try:
            from openpyxl import load_workbook
        except ImportError as exc: raise ValueError("XLSX support is unavailable") from exc
        try: wb=load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as exc: raise ValueError("invalid XLSX document") from exc
        parts=[]
        try:
            if len(wb.worksheets) > MAX_XLSX_SHEETS:
                raise ValueError("XLSX exceeds 30 sheet limit")
            for ws in wb.worksheets:
                if (ws.max_column or 0) > MAX_XLSX_COLUMNS:
                    raise ValueError("XLSX exceeds 200 column limit")
                first_row = next(ws.iter_rows(max_row=1, max_col=MAX_XLSX_COLUMNS), None)
                if first_row is None: raise ValueError(f"sheet {ws.title} is empty")
                headers=[str(x.value or "").strip() for x in first_row]
                required={"report_date","area","asset_tag","work_description","quantity","unit","quantity_kind"}
                if not required.issubset(set(headers)): raise ValueError(f"sheet {ws.title} missing required columns")
                for ri,row in enumerate(ws.iter_rows(min_row=2, max_col=MAX_XLSX_COLUMNS),2):
                    if ri>MAX_XLSX_ROWS+1: raise ValueError("XLSX exceeds 5,000 row limit")
                    vals=[x.value for x in row]; text=" | ".join(f"{h}: {v}" for h,v in zip(headers,vals) if v is not None)
                    if text.strip(): add_part(parts, f"sheet:{ws.title}/row:{ri}",text)
        finally:
            wb.close()
        if not parts: raise ValueError("XLSX contains no report rows")
        return parts,warnings,"xlsx"
    raise ValueError("unsupported report format")

def ingest_report(db: Session, project: Project, raw: bytes, *, filename: str|None=None, mime: str="text/plain", report_date: date|None=None, report_date_evidence: str|None=None, source_label: str|None=None, uploader_id=None, mapping=None, ingestion_metadata=None, parsed_parts=None):
    if not raw: raise ValueError("report is empty")
    if len(raw)>MAX_BYTES: raise ValueError("report exceeds 10 MB limit")
    name=_filename(filename); digest=hashlib.sha256(raw).hexdigest()
    # The existing unique constraint covers project and content_hash. Include the
    # work date in that identity while preserving the raw checksum on FileRecord.
    metadata = dict(ingestion_metadata or {})
    if mapping is not None:
        from .report_mapping import mapped_rows
        fragments, errors, mapped_metadata = mapped_rows(raw, mapping)
        if errors:
            raise ValueError("XLSX_ROW_ERRORS:" + json.dumps(errors))
        if not fragments:
            raise ValueError("XLSX contains no report rows")
        metadata.update(mapped_metadata)
        warnings, kind = [], "xlsx"
    elif parsed_parts is not None:
        fragments, warnings, kind = parsed_parts, [], "transcription"
    else:
        fragments,warnings,kind=parse_bytes(raw,name,mime)
    identity_context = json.dumps(metadata, sort_keys=True, default=str) if metadata else ""
    identity=hashlib.sha256(f"{digest}:{report_date.isoformat() if report_date else ''}:{identity_context}".encode()).hexdigest() if metadata else hashlib.sha256(f"{digest}:{report_date.isoformat() if report_date else ''}".encode()).hexdigest()
    existing=db.scalar(select(Report).where(Report.project_id==project.id, or_(
        Report.content_hash==identity,
        and_(Report.content_hash==digest, Report.report_date==report_date) if not metadata else False,
    )))
    if existing: return existing, True
    if sum(len(text) for _, text in fragments)>MAX_EXTRACTED_CHARS: raise ValueError("extracted report text exceeds limit")
    storage_key=f"{project.id}/{uuid.uuid4().hex}-{name}"
    _store_original(storage_key, raw)
    try:
        record=FileRecord(project_id=project.id, original_filename=name, storage_key=storage_key, sha256=digest, mime_type=mime, size=len(raw), source_kind=kind, uploader_id=uploader_id)
        db.add(record); db.flush()
        report=Report(project_id=project.id,file_id=record.id,report_date=report_date,report_date_evidence=report_date_evidence,source_label=source_label,content_hash=identity,parsing_warnings=json.dumps(warnings),ingestion_metadata=metadata or None)
        db.add(report); db.flush()
        for ordinal,(locator,text) in enumerate(fragments,1):
            db.add(Fragment(report_id=report.id,ordinal=ordinal,locator=locator,original_text=text,normalised_text=_normalise(text),ocr_status="NEEDS_TRANSCRIPTION" if locator in {w.split(":",1)[1] for w in warnings if w.startswith("NEEDS_TRANSCRIPTION:")} else "NOT_REQUIRED"))
        return report, False
    except Exception:
        _remove_original(storage_key)
        raise
