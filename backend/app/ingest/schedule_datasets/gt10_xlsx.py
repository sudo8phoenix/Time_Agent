"""GT10 workbook reader; preserves source cells and rejects broken references."""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook


def _value(value):
    if isinstance(value, (date, datetime)):
        return {"raw": value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat(),
                "cell_type": "datetime", "precision": "source_datetime" if isinstance(value, datetime) else "source_date",
                "timezone_basis": "unspecified"}
    return value


def _rows(book, sheet_name: str, header_rows: int) -> Iterator[dict]:
    sheet = book[sheet_name]
    rows = sheet.iter_rows()
    headers = [cell.value for cell in next(rows)]
    for _ in range(header_rows - 1):
        next(rows)
    for row in rows:
        if not any(cell.value is not None for cell in row):
            continue
        if any(cell.data_type == "e" for cell in row):
            raise ValueError(f"spreadsheet error in {sheet_name}!{row[0].row}")
        yield {"locator": f"{sheet_name}!{row[0].row}",
               "cells": {str(key): _value(cell.value) for key, cell in zip(headers, row)}}


def read_gt10(path: str | Path, *, version: str = "export") -> dict:
    """Return exact schedule rows with locators; no inferred actuals or mappings."""
    if version not in {"export", "input"}:
        raise ValueError("version must be export or input")
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        task_sheet, edge_sheet = (("TASK", "TASKPRED") if version == "export"
                                  else ("ACTIVITES", "RELATIONSHIP"))
        offset = 2 if version == "export" else 1
        tasks = list(_rows(book, task_sheet, offset))
        edges = list(_rows(book, edge_sheet, offset))
        task_key = "task_code" if version == "export" else "Activity ID"
        pred_key = "pred_task_id" if version == "export" else "Predecessor Activity ID"
        succ_key = "task_id" if version == "export" else "Successor Activity ID"
        ids = [str(row["cells"].get(task_key) or "") for row in tasks]
        if len(ids) != len(set(ids)) or any(not value for value in ids):
            raise ValueError("duplicate or blank task ID")
        known = set(ids)
        issues = []
        for row in edges:
            pred, succ = (str(row["cells"].get(key) or "") for key in (pred_key, succ_key))
            if pred not in known or succ not in known:
                issues.append({"code": "unresolved_link", "locator": row["locator"],
                               "predecessor": pred, "successor": succ})
        if issues:
            raise ValueError(f"unresolved links: {issues[:5]}")
        return {"source": "GT10", "version": version, "project_id": "GT10BLDG",
                "task_count": len(tasks), "relationship_count": len(edges),
                "tasks": tasks, "relationships": edges,
                "wbs": list(_rows(book, "WBS", 1)) if version == "input" else None,
                "calendar_status": "incomplete_unreviewed", "timezone_basis": "unspecified",
                "review_status": "unreviewed"}
    finally:
        book.close()
