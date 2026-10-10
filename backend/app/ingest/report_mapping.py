"""Explicit, revisioned report mappings; raw cell evidence is never discarded."""
from datetime import date, datetime, time
import hashlib
import io
import json
import re
import zipfile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles.numbers import is_datetime
from pydantic import BaseModel, ConfigDict, Field, model_validator

FIELDS = ('work_description', 'report_date', 'area', 'asset_tag', 'discipline',
          'event_kind', 'actual_start', 'actual_finish', 'quantity', 'unit', 'quantity_kind')


class MappingSpec(BaseModel):
    model_config = ConfigDict(extra='forbid')
    sheet: str = Field(min_length=1, max_length=100)
    header_row: int = Field(ge=1, le=100)
    columns: dict[str, int]
    context: dict[str, str] = Field(default_factory=dict)
    date_format: str | None = None

    @model_validator(mode='after')
    def validate_mapping(self):
        if not self.columns or set(self.columns) - set(FIELDS):
            raise ValueError('unknown or empty canonical column mapping')
        if any(type(v) is not int or v < 1 or v > 200 for v in self.columns.values()):
            raise ValueError('column indices must be 1–200')
        if len(set(self.columns.values())) != len(self.columns):
            raise ValueError('a source column can only map to one field')
        if set(self.context) - set(FIELDS) or any(len(v) > 2000 for v in self.context.values()):
            raise ValueError('invalid file context')
        if 'work_description' not in self.columns:
            raise ValueError('map a work description column')
        for key in ('report_date', 'actual_start', 'actual_finish'):
            if self.context.get(key):
                try:
                    if key == 'report_date':
                        date.fromisoformat(self.context[key])
                    else:
                        datetime.fromisoformat(self.context[key])
                except ValueError as exc:
                    raise ValueError('file date context must be ISO and explicit') from exc
        if self.date_format not in (None, '%d/%m/%Y', '%m/%d/%Y'):
            raise ValueError('choose D/M/Y or M/D/Y explicitly')
        return self


def workbook(raw):
    from .reports import MAX_BYTES, MAX_ZIP_EXPANDED_BYTES
    if len(raw) > MAX_BYTES:
        raise ValueError('report exceeds 10 MB limit')
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_ZIP_EXPANDED_BYTES:
                raise ValueError('document expands beyond 50 MB limit')
        return load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    except (zipfile.BadZipFile, KeyError) as exc:
        raise ValueError('invalid XLSX document') from exc


def cell_value(value):
    return value.isoformat() if isinstance(value, (date, datetime, time)) else str(value) if value is not None else ''


def signature(sheet, row, headers):
    return hashlib.sha256(json.dumps([sheet, row, headers], ensure_ascii=False).encode()).hexdigest()


def preview(raw):
    wb = workbook(raw)
    try:
        result = []
        for ws in wb.worksheets:
            if len(result) >= 30:
                raise ValueError('too many sheets to preview')
            rows = []
            for ri, row in enumerate(ws.iter_rows(max_row=20, max_col=min(ws.max_column or 1, 200)), 1):
                rows.append([{'cell': f"{get_column_letter(ci)}{ri}", 'value': cell_value(c.value)[:500]} for ci, c in enumerate(row, 1)])
            result.append({'name': ws.title, 'rows': rows})
        return result
    finally:
        wb.close()


def mapped_rows(raw, spec: MappingSpec):
    from .reports import MAX_XLSX_ROWS, MAX_EXTRACTED_CHARS
    wb = workbook(raw)
    try:
        if spec.sheet not in wb.sheetnames:
            raise ValueError('mapped sheet is missing')
        ws = wb[spec.sheet]
        if (ws.max_column or 0) > 200:
            raise ValueError('report exceeds 200 columns')
        header = next(ws.iter_rows(min_row=spec.header_row, max_row=spec.header_row), ())
        headers = [cell_value(c.value) for c in header]
        if not headers or any(v > len(headers) for v in spec.columns.values()):
            raise ValueError('mapped column is missing')
        sig = signature(spec.sheet, spec.header_row, headers)
        parts, errors, evidence = [], [], []
        total = 0
        for ri, row in enumerate(ws.iter_rows(min_row=spec.header_row + 1), spec.header_row + 1):
            if ri - spec.header_row > MAX_XLSX_ROWS:
                raise ValueError('XLSX exceeds 5,000 row limit')
            if not any(c.value is not None for c in row):
                continue
            values = dict(spec.context)
            cells = [{'cell': f'{get_column_letter(i + 1)}{ri}', 'header': headers[i] if i < len(headers) else '',
                      'value': cell_value(c.value), 'type': getattr(c, 'data_type', 'n'), 'number_format': getattr(c, 'number_format', None)} for i, c in enumerate(row)]
            row_errors = []
            for key, col in spec.columns.items():
                c = row[col - 1]
                value = cell_value(c.value)
                coordinate = f'{get_column_letter(col)}{ri}'
                if getattr(c, 'data_type', None) == 'f':
                    row_errors.append(f'{coordinate}: formula requires an explicit value')
                if value and key in ('report_date', 'actual_start', 'actual_finish'):
                    if isinstance(c.value, datetime) and (key == 'report_date' or is_datetime(c.number_format) == 'date'):
                        value = c.value.date().isoformat()
                    elif not isinstance(c.value, (date, datetime, time)):
                        try:
                            if key == 'report_date':
                                value = (datetime.strptime(value, spec.date_format).date().isoformat()
                                         if spec.date_format else date.fromisoformat(value).isoformat())
                            elif re.fullmatch(r'\d{2}:\d{2}(:\d{2})?', value):
                                time.fromisoformat(value)
                            elif spec.date_format and '/' in value:
                                value = datetime.strptime(value, spec.date_format).date().isoformat()
                            else:
                                datetime.fromisoformat(value)
                        except ValueError:
                            row_errors.append(f'{coordinate}: ambiguous/invalid date; choose date format or ISO')
                if value:
                    values[key] = value
            if not values.get('work_description', '').strip():
                row_errors.append('missing work description')
            locator = f'sheet:{ws.title}/row:{ri}'
            evidence.append({'locator': locator, 'cells': cells, 'context': spec.context, 'errors': row_errors})
            if row_errors:
                errors.append({'locator': locator, 'errors': row_errors})
                continue
            text = ' | '.join(f'{k}: {v}' for k, v in values.items())
            total += len(text) + sum(len(c['value']) for c in cells)
            if total > MAX_EXTRACTED_CHARS:
                raise ValueError('extracted report text exceeds limit')
            parts.append((locator, text))
        return parts, errors, {'header_signature': sig, 'headers': headers,
                               'mapping': spec.model_dump(), 'rows': evidence}
    finally:
        wb.close()
