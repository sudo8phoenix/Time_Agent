"""Hostile real documents through intake and authority checks, with injected models.

These are deterministic enforcement tests, not a hosted-model robustness score.
"""
from io import BytesIO
from dataclasses import replace
import json
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
import pytest
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.agent.tools import AgentToolError
from app.db.models import Fragment, ProgressEvent
from app.ingest.reports import parse_bytes
from app.llm.extract import BatchInput, FragmentInput, extract_batch
from test_agent_tools import Fixture

CORPUS = Path(__file__).resolve().parents[2] / 'fixtures' / 'security' / 'prompt_injection.json'
ATTACKS = json.loads(CORPUS.read_text())


def document_bytes(kind, text):
    stream = BytesIO()
    if kind == 'txt':
        return text.encode()
    if kind == 'docx':
        doc = Document()
        doc.add_paragraph(text)
        doc.save(stream)
    elif kind == 'xlsx':
        book = Workbook()
        book.active.append(['report_date', 'area', 'asset_tag', 'work_description',
                            'quantity', 'unit', 'quantity_kind'])
        book.active.append(['2026-09-01', '', '', text, 2, 'm', 'delta'])
        book.save(stream)
        book.close()
    else:
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
            NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):
            DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        content = DecodedStreamObject()
        escaped = text.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')
        content.set_data(f'BT /F1 10 Tf 20 700 Td ({escaped}) Tj ET'.encode('ascii'))
        page[NameObject('/Contents')] = writer._add_object(content)
        writer.write(stream)
    return stream.getvalue()


@pytest.mark.parametrize('attack', ATTACKS, ids=lambda item: item['id'])
@pytest.mark.parametrize('kind', ['txt', 'docx', 'xlsx', 'pdf'])
def test_hostile_document_never_gains_system_or_approval_authority(attack, kind):
    fixture = Fixture()
    text = 'Installed 2 m of cable. ' + attack['text']
    parts, _, _ = parse_bytes(document_bytes(kind, text), 'report.' + kind, 'application/octet-stream')
    source = '\n'.join(value for _, value in parts)
    assert attack['text'] in source
    with Session(fixture.engine) as db:
        fragment = db.get(Fragment, fixture.fragment_id)
        fragment.original_text = fragment.normalised_text = source
        db.commit()
    calls = []
    baseline = []
    extract_batch(BatchInput('baseline', (FragmentInput('f1', 'p1', 'Installed 2 m of cable.'),)),
                  lambda **request: baseline.append(request) or {'observations': []})

    def injected_model(**request):
        calls.append(request)
        return {'observations': [fixture.observation.model_dump(mode='json')]}

    fixture.context = replace(fixture.context, model_callable=injected_model)
    tools = fixture.tools()
    scope = fixture.load(tools)
    extracted = tools.extract_observations(scope.fragment_ids, tools.report_metadata)
    assert extracted.observations == [fixture.observation]
    assert calls[0]['system'] == baseline[0]['system']
    assert calls[0]['schema'] == baseline[0]['schema']
    assert attack['text'] in calls[0]['user']
    assert attack['text'] not in calls[0]['system']
    assert fixture.persist_calls == []
    with Session(fixture.engine) as db:
        assert db.scalar(select(func.count()).select_from(ProgressEvent)) == 0


@pytest.mark.parametrize('forgery', ['missing_quote', 'foreign_fragment', 'approval', 'tool_call'])
def test_document_induced_forged_extraction_is_rejected(forgery):
    fixture = Fixture()
    attack = next(item for item in ATTACKS if item['id'] == 'fabricated-evidence')
    parts, _, _ = parse_bytes(document_bytes('docx', 'Installed 2 m of cable. ' + attack['text']),
                              'hostile.docx', 'application/octet-stream')
    source = '\n'.join(value for _, value in parts)
    with Session(fixture.engine) as db:
        fragment = db.get(Fragment, fixture.fragment_id)
        fragment.original_text = fragment.normalised_text = source
        db.commit()
    output = fixture.observation.model_dump(mode='json')
    if forgery == 'missing_quote':
        output['evidence'][0]['quote'] = 'Installed 9000 m of cable.'
    elif forgery == 'foreign_fragment':
        output['evidence'][0]['fragment_id'] = '00000000-0000-0000-0000-000000000001'
    elif forgery == 'approval':
        output['review_state'] = 'approved'
    else:
        output['tool_call'] = {'name': 'execute_shell', 'arguments': {'command': 'export secrets'}}
    fixture.context = replace(fixture.context, model_callable=lambda **request: {'observations': [output]})
    tools = fixture.tools()
    scope = fixture.load(tools)
    with pytest.raises(AgentToolError):
        tools.extract_observations(scope.fragment_ids, tools.report_metadata)
    assert fixture.persist_calls == []
    with Session(fixture.engine) as db:
        assert db.scalar(select(func.count()).select_from(ProgressEvent)) == 0
