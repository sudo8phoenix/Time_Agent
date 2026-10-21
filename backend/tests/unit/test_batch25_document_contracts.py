"""Actual parser/ledger boundaries and auditable, unadjudicated corpus contracts."""
from copy import deepcopy
from datetime import date
from io import BytesIO
import json
from pathlib import Path

from openpyxl import Workbook
import pytest

from app.ingest import reports
from app.ingest.report_mapping import MappingSpec, _mapped_rows_local
from app.llm.extract import FragmentInput, ExtractionError, build_batches
from app.progress.ledger import ActivityBaseline, ProgressEvent, recompute_progress
from evaluation.development_contract import score_case
from evaluation.document_corpus_contract import audit, DIMENSIONS, STRESS_SLICES
from evaluation.source_units import source_units

ROOT = Path(__file__).resolve().parents[3]
CORPUS = json.loads((ROOT / 'evaluation/fixtures/regression/document-corpus.json').read_text())


def workbook(rows, *, merged=False, titles=('Diary',)):
    book = Workbook()
    for i, title in enumerate(titles):
        sheet = book.active if i == 0 else book.create_sheet()
        sheet.title = title
        for row in rows:
            sheet.append(row)
        if merged:
            sheet.merge_cells('B2:B3')
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def spec():
    return MappingSpec(sheet='Diary', header_row=1, columns={'work_description': 1, 'area': 2})


def test_qual11_incompatible_cumulative_assertions_both_remain_conflicts_with_no_percent():
    events = [ProgressEvent('source-A', date(2026, 10, 21), 4, 'm', 'cumulative'),
              ProgressEvent('source-B', date(2026, 10, 21), 7, 'm', 'cumulative')]
    original = deepcopy(events)
    baseline = ActivityBaseline(100, 'm', 0)
    for order in [events, list(reversed(events))]:
        result = recompute_progress(baseline, order)
        assert result.conflicts and 'conflicting cumulative totals' in result.conflicts[0]
        assert result.physical_percent is None and result.applied_event_ids == []
        assert result.completed_quantity == 0
    assert events == original


def test_qual12_repeated_headers_are_review_errors_with_complete_raw_cell_evidence():
    rows = [['Work', 'Area'], ['Erected 2 spools\nWelded 3 joints', 'Zone A'],
            ['Work', 'Area'], ['Delivered 5 spools*', 'Zone B'], ['* delivery is not installation', None]]
    parts, errors, metadata = _mapped_rows_local(workbook(rows), spec())
    assert [locator for locator, _ in parts] == ['sheet:Diary/row:2', 'sheet:Diary/row:4', 'sheet:Diary/row:5']
    assert '\nWelded 3 joints' in parts[0][1]
    assert errors[0]['locator'] == 'sheet:Diary/row:3'
    assert 'repeated header' in errors[0]['errors'][0]
    assert len(metadata['rows']) == 4
    assert [cell['value'] for cell in metadata['rows'][1]['cells']] == rows[0]
    assert metadata['headers'] == rows[0]
    assert 'delivery is not installation' in parts[-1][1]


def test_qual12_merged_cell_does_not_invent_second_row_scope():
    parts, errors, metadata = _mapped_rows_local(workbook([
        ['Work', 'Area'], ['Installed pipe', 'Zone A'], ['Welded joints', None]], merged=True), spec())
    assert not errors
    assert 'area: Zone A' in parts[0][1] and 'area:' not in parts[1][1]
    assert metadata['rows'][1]['cells'][1]['value'] == ''
    assert metadata['rows'][1]['cells'][1]['cell'] == 'B3'


@pytest.mark.parametrize('text', ['ढलाई 2 m3; cure reqd', 'бетон 3 м³', 'صب 2 m3',
                                  'béton posé à Zone-A', '焊接 3 joints', 'wldg 3 jts / P-001'])
def test_qual13_utf8_parser_and_whitespace_normalizer_keep_scripts_accents_and_abbreviations(text):
    raw = ('\ufeff' + text + '\n\n' + text).encode('utf-8')
    parts, warnings, kind = reports._parse_bytes_local(raw, 'report.txt', 'text/plain')
    assert parts == [('paragraph:1', text), ('paragraph:2', text)]
    assert not warnings and kind == 'text'
    assert reports._normalise(' \t' + text + '\n ') == text
    units = source_units(text)
    assert ''.join(unit['focus'] for unit in units) == text


@pytest.mark.parametrize('count', [8, 9, 30])
def test_qual14_long_reports_keep_first_middle_last_facts_and_all_fragment_identities(count):
    paragraphs = [f'FACT-{i:03}: '+('context '*1000)+f' END-{i:03}' for i in range(count)]
    parts, _, _ = reports._parse_bytes_local('\n\n'.join(paragraphs).encode(), 'long.txt', 'text/plain')
    fragments = [FragmentInput(f'f{i}', locator, text, i) for i, (locator, text) in enumerate(parts)]
    batches = build_batches(fragments)
    flattened = [fragment for batch in batches for fragment in batch.fragments]
    assert flattened == fragments and [fragment.text for fragment in flattened] == paragraphs
    assert flattened[0].text.startswith('FACT-000')
    assert flattened[count//2].text.startswith(f'FACT-{count//2:03}')
    assert flattened[-1].text.endswith(f'END-{count-1:03}')
    assert all(sum(len(f.text) for f in batch.fragments) <= 12000 for batch in batches)


def test_qual14_oversized_single_fragment_fails_explicitly_without_partial_batch_result():
    fragments = [FragmentInput('first', 'p1', 'first'), FragmentInput('large', 'p2', 'x'*12001),
                 FragmentInput('last', 'p3', 'last')]
    with pytest.raises(ExtractionError, match='large exceeds batch limit'):
        build_batches(fragments)


@pytest.mark.parametrize('raw,message', [(b'', 'empty'), (b' \n\t', 'empty'), (b'\xff', 'UTF-8')])
def test_qual15_partial_empty_unreadable_text_is_rejected_before_extraction(raw, message):
    with pytest.raises(ValueError, match=message):
        reports._parse_bytes_local(raw, 'empty.txt', 'text/plain')


def test_doc05_sheet_order_and_raw_row_limits_are_enforced_without_truncation(monkeypatch):
    headers = ['report_date', 'area', 'asset_tag', 'work_description', 'quantity', 'unit', 'quantity_kind']
    rows = [headers, ['2026-10-21', 'A', '001', 'first', 1, 'm', 'delta'],
            ['2026-10-21', 'B', '002', 'last', 2, 'm', 'delta']]
    raw = workbook(rows, titles=('Z-last-name', 'A-first-name'))
    parts, _, kind = reports._parse_bytes_local(raw, 'sheets.xlsx', 'application/octet-stream')
    assert kind == 'xlsx'
    assert [locator for locator, _ in parts] == ['sheet:Z-last-name/row:2', 'sheet:Z-last-name/row:3',
                                                'sheet:A-first-name/row:2', 'sheet:A-first-name/row:3']
    monkeypatch.setattr(reports, 'MAX_XLSX_ROWS', 1)
    with pytest.raises(ValueError, match='row limit'):
        reports._parse_bytes_local(raw, 'sheets.xlsx', 'application/octet-stream')


def test_data05_all_source_units_have_explicit_atomic_or_no_event_dispositions():
    result = audit(CORPUS)
    assert result['cases'] == 12 and result['annotated_events'] == 12
    assert result['semantic_labels_verified'] is False
    events = CORPUS['cases'][0]['units'][0]['events']
    contract = {'alternatives': [{'id': 'all-atomic', 'events': events}]}
    assert score_case(contract, events, completed=True)['fields_exact_any_alternative']
    assert not score_case(contract, events[:1], completed=True)['fields_exact_any_alternative']
    for mutation in ['drop_unit', 'wrong_span', 'missing_evidence', 'wrong_source', 'unexplained_empty']:
        corpus = deepcopy(CORPUS)
        case = corpus['cases'][0]
        if mutation == 'drop_unit':
            case['units'].pop()
        elif mutation == 'wrong_span':
            case['units'][0]['start'] += 1
        elif mutation == 'missing_evidence':
            case['units'][0]['events'][0]['evidence'] = []
        elif mutation == 'wrong_source':
            case['units'][0]['events'][0]['evidence'][0]['fragment_id'] = 'other'
        else:
            case['units'][0].update(disposition='no_event', events=[], reason='')
        with pytest.raises(ValueError):
            audit(corpus)


def test_data06_coherent_ambiguity_alternatives_reject_hybrid_and_omission():
    case = next(case for case in CORPUS['cases'] if 'ambiguity' in case['slices'])
    for alternative in case['alternatives']:
        result = score_case(case, alternative['events'], completed=True)
        assert result['accepted_alternative_ids'] == [alternative['id']]
        assert result['semantic_evidence_correct'] is None
    hybrid = deepcopy(case['alternatives'][0]['events'])
    hybrid[0]['work_date'] = case['alternatives'][1]['events'][0]['work_date']
    assert not score_case(case, hybrid, completed=True)['fields_exact_any_alternative']
    assert not score_case(case, [], completed=True)['fields_exact_any_alternative']
    broken = deepcopy(CORPUS)
    next(case for case in broken['cases'] if 'ambiguity' in case['slices'])['alternatives'] = []
    with pytest.raises(ValueError, match='nonempty alternative'):
        audit(broken)


@pytest.mark.parametrize('dimension', DIMENSIONS)
def test_data09_diversity_inventory_is_counted_and_missing_dimensions_fail(dimension):
    result = audit(CORPUS)
    assert len(result['dimensions'][dimension]) >= 3
    assert sum(result['dimensions'][dimension].values()) == 12
    corpus = deepcopy(CORPUS)
    del corpus['cases'][0][dimension]
    with pytest.raises(ValueError, match='diversity dimension'):
        audit(corpus)


def test_data10_every_difficult_slice_is_counted_without_claiming_natural_prevalence():
    result = audit(CORPUS)
    assert set(result['slices']) == STRESS_SLICES
    assert result['missing_stress_slices'] == [] and result['prevalence_estimate'] is False
    assert result['slices']['ordinary'] == 5
    removed = deepcopy(CORPUS)
    removed['cases'] = [case for case in removed['cases'] if 'wrong_project' not in case['slices']]
    assert audit(removed)['missing_stress_slices'] == ['wrong_project']
