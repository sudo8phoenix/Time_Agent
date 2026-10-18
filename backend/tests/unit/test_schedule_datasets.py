"""Schedule adapter contracts using self-contained synthetic source files."""
from datetime import datetime
import json

import pytest

from app.ingest.schedule_datasets.cambridge import read_cambridge
from app.ingest.schedule_datasets.gt10_xlsx import read_gt10

def test_gt10_export_and_input_keep_identity_and_raw_dates(tmp_path):
    from openpyxl import Workbook

    # Export has technical and display headers; input has a single header.
    # Mismatched source names must survive rather than being silently repaired.
    tasks = [("SYN-01", "Excavation", datetime(2026, 10, 1, 8, 30)),
             ("SYN-F8", "De-shuttering - F9", datetime(2026, 10, 2, 0, 0))]
    paths = {}
    for version in ("export", "input"):
        book = Workbook()
        task = book.active
        task.title = "TASK" if version == "export" else "ACTIVITES"
        task.append(["task_code", "task_name", "start_date"] if version == "export"
                    else ["Activity ID", "Activity Name", "Start"])
        if version == "export":
            task.append(["Activity ID", "Activity Name", "Start"])
        for row in tasks:
            task.append(row)
        edge = book.create_sheet("TASKPRED" if version == "export" else "RELATIONSHIP")
        edge.append(["pred_task_id", "task_id"] if version == "export"
                    else ["Predecessor Activity ID", "Successor Activity ID"])
        if version == "export":
            edge.append(["Predecessor", "Successor"])
        edge.append(["SYN-01", "SYN-F8"])
        if version == "input":
            wbs = book.create_sheet("WBS")
            wbs.append(["WBS Code"])
            wbs.append(["SYNTHETIC"])
        paths[version] = tmp_path / f"{version}.xlsx"
        book.save(paths[version])
        book.close()

    export = read_gt10(paths["export"])
    source = read_gt10(paths["input"], version="input")
    assert (export["task_count"], export["relationship_count"]) == (2, 1)
    assert (source["task_count"], source["relationship_count"]) == (2, 1)
    assert {row["cells"]["task_code"] for row in export["tasks"]} == {
        row["cells"]["Activity ID"] for row in source["tasks"]
    }
    assert export["tasks"][0]["locator"] == "TASK!3"
    assert source["tasks"][0]["locator"] == "ACTIVITES!2"
    assert source["wbs"][0]["cells"]["WBS Code"] == "SYNTHETIC"
    assert export["relationships"][0]["cells"] == {"pred_task_id": "SYN-01", "task_id": "SYN-F8"}
    assert source["relationships"][0]["cells"] == {
        "Predecessor Activity ID": "SYN-01", "Successor Activity ID": "SYN-F8"}
    f8 = next(row for row in export["tasks"] if row["cells"]["task_code"] == "SYN-F8")
    assert f8["cells"]["task_name"] == "De-shuttering - F9"
    assert f8["cells"]["start_date"]["precision"] == "source_datetime"
    assert f8["cells"]["start_date"]["timezone_basis"] == "unspecified"
    assert f8["cells"]["start_date"]["raw"] == "2026-10-02 00:00:00"
    assert source["tasks"][1]["cells"]["Start"] == f8["cells"]["start_date"]
    assert not any(key.startswith("actual") for key in f8["cells"])


def test_gt10_rejects_unresolved_link(tmp_path):
    from openpyxl import Workbook
    book = Workbook()
    task = book.active
    task.title = "TASK"
    task.append(["task_code"])
    task.append(["Activity ID"])
    task.append(["A"])
    edge = book.create_sheet("TASKPRED")
    edge.append(["pred_task_id", "task_id"])
    edge.append(["Predecessor", "Successor"])
    edge.append(["A", "MISSING"])
    path = tmp_path / "broken.xlsx"
    book.save(path)
    with pytest.raises(ValueError, match="unresolved links"):
        read_gt10(path)


def test_cambridge_stream_preserves_composite_id_and_raw_values(tmp_path):
    raw = {"Task_ID": "T1", "Actual_Start_Date": "2026-10-01 00:00",
           "Actual_Finish_Date": None, "Activity_Relationships": "null",
           "unknown_source_field": {"value": "retain unchanged"}}
    path = tmp_path / "synthetic-cambridge.json"
    path.write_text(json.dumps({"Projects": {
        "SYN-A": {"Project_ID": "SYN-A", "Activities": {"T1": raw}},
        "SYN-B": {"Project_ID": "SYN-B", "Activities": {"T1": raw}},
    }}))
    items = list(read_cambridge(path))
    assert [item["project_id"] for item in items] == ["SYN-A", "SYN-B"]
    item = items[0]
    assert len(item["activities"]) == 1
    activity = item["activities"][0]
    assert activity["identity"] == ["cambridge", "SYN-A", "T1"]
    assert items[1]["activities"][0]["identity"] == ["cambridge", "SYN-B", "T1"]
    assert activity["raw"] == raw
    assert activity["locator"] == '$.Projects["SYN-A"].Activities["T1"]'
    assert activity["relationships"] == {}
    assert activity["timestamp_basis"] == "source_string_timezone_unspecified"
    assert list(read_cambridge(path, project_id="SYN-B")) == [items[1]]
