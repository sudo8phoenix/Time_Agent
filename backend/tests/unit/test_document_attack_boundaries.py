"""Bounded, offline adversarial document intake tests; no inference adapters."""
from io import BytesIO
import stat
import zipfile

from docx import Document
from openpyxl import Workbook
import pytest

from app.ingest import reports
from app.ingest.report_mapping import workbook


def archive_bytes(entries):
    stream = BytesIO()
    with zipfile.ZipFile(stream, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in entries:
            archive.writestr(name, content)
    return stream.getvalue()


@pytest.mark.parametrize('entry', ['../escape.xml', '/absolute.xml', 'word/../../escape.xml',
                                  'C:/escape.xml', 'word\\escape.xml'])
@pytest.mark.parametrize('entrypoint', ['docx', 'xlsx', 'mapping'])
def test_unsafe_archive_members_rejected_before_document_library(entry, entrypoint, monkeypatch):
    raw = archive_bytes([(entry, 'hostile')])
    monkeypatch.setattr('docx.Document', lambda *a: pytest.fail('DOCX library reached'))
    monkeypatch.setattr('openpyxl.load_workbook', lambda *a, **k: pytest.fail('XLSX library reached'))
    monkeypatch.setattr('app.ingest.report_mapping.load_workbook', lambda *a, **k: pytest.fail('mapping library reached'))
    with pytest.raises(ValueError, match='unsafe or ambiguous'):
        if entrypoint == 'mapping':
            workbook(raw)
        else:
            reports._parse_bytes_local(raw, 'attack.' + entrypoint, 'application/octet-stream')


def test_duplicate_member_and_symbolic_link_are_rejected():
    with pytest.warns(UserWarning, match='Duplicate name'):
        duplicate = archive_bytes([('word/document.xml', 'one'), ('word/document.xml', 'two')])
    with pytest.raises(ValueError, match='unsafe or ambiguous'):
        reports.validate_document_archive(duplicate)
    link = zipfile.ZipInfo('word/link')
    link.create_system = 3
    link.external_attr = (stat.S_IFLNK | 0o777) << 16
    with pytest.raises(ValueError, match='unsafe or ambiguous'):
        reports.validate_document_archive(archive_bytes([(link, '/etc/passwd')]))


def test_archive_member_count_is_bounded():
    raw = archive_bytes([(f'part{i}.xml', '') for i in range(reports.MAX_ZIP_MEMBERS + 1)])
    with pytest.raises(ValueError, match='too many archive members'):
        reports.validate_document_archive(raw)


def test_compressed_document_expansion_is_bounded(monkeypatch):
    monkeypatch.setattr(reports, 'MAX_ZIP_EXPANDED_BYTES', 1024)
    raw = archive_bytes([('word/document.xml', 'A' * 1025)])
    assert len(raw) < 1024
    with pytest.raises(ValueError, match='expands beyond'):
        reports.validate_document_archive(raw)


@pytest.mark.parametrize('suffix', ['txt', 'pdf', 'docx', 'xlsx'])
def test_raw_byte_limit_applies_at_parser_boundary(suffix, monkeypatch):
    monkeypatch.setattr(reports, 'MAX_BYTES', 10)
    with pytest.raises(ValueError, match='10 MB'):
        reports._parse_bytes_local(b'x' * 11, 'large.' + suffix, 'application/octet-stream')


def test_docx_interleaved_paragraph_and_table_preserve_source_order():
    document = Document()
    document.add_paragraph('Before table')
    document.add_table(rows=1, cols=1).cell(0, 0).text = 'Middle table'
    document.add_paragraph('After table')
    raw = BytesIO()
    document.save(raw)
    parts, warnings, kind = reports._parse_bytes_local(raw.getvalue(), 'report.docx', 'application/octet-stream')
    assert parts == [('paragraph:1', 'Before table'), ('table:1/row:1', 'Middle table'),
                     ('paragraph:2', 'After table')]
    assert warnings == [] and kind == 'docx'


def spreadsheet(*, columns=7, sheets=1):
    book = Workbook()
    headers = ['report_date', 'area', 'asset_tag', 'work_description', 'quantity', 'unit', 'quantity_kind']
    for index in range(sheets):
        sheet = book.active if index == 0 else book.create_sheet()
        sheet.append(headers + ['extra'] * (columns - len(headers)))
        sheet.append(['2026-10-20', 'A', 'P-1', 'Installed pipe', 2, 'm', 'delta'])
    raw = BytesIO()
    book.save(raw)
    book.close()
    return raw.getvalue()


@pytest.mark.parametrize('columns,sheets,message', [(201, 1, '200 column'), (7, 31, '30 sheet')])
def test_spreadsheet_dimensions_are_bounded_and_workbook_closed(columns, sheets, message, monkeypatch):
    import openpyxl
    original = openpyxl.load_workbook
    closed = []

    def load(*args, **kwargs):
        book = original(*args, **kwargs)
        original_close = book.close
        def close():
            closed.append(True)
            original_close()
        book.close = close
        return book

    monkeypatch.setattr(openpyxl, 'load_workbook', load)
    with pytest.raises(ValueError, match=message):
        reports._parse_bytes_local(spreadsheet(columns=columns, sheets=sheets), 'report.xlsx', 'application/octet-stream')
    assert closed == [True]


def test_spreadsheet_sheet_order_and_row_locators_preserved():
    parts, warnings, kind = reports._parse_bytes_local(spreadsheet(sheets=2), 'report.xlsx', 'application/octet-stream')
    assert [locator for locator, text in parts] == ['sheet:Sheet/row:2', 'sheet:Sheet1/row:2']
    assert warnings == [] and kind == 'xlsx'


def test_image_upload_is_explicitly_unsupported_by_text_parser():
    with pytest.raises(ValueError, match='unsupported report format'):
        reports._parse_bytes_local(b'\x89PNG\r\n', 'scan.png', 'image/png')


def pdf_bytes(texts, *, encrypted=False):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=612, height=792)
        if text:
            font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                                     NameObject('/Subtype'): NameObject('/Type1'),
                                     NameObject('/BaseFont'): NameObject('/Helvetica')})
            page[NameObject('/Resources')] = DictionaryObject({
                NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
            content = DecodedStreamObject()
            content.set_data(f'BT /F1 12 Tf 50 700 Td ({text}) Tj ET'.encode())
            page[NameObject('/Contents')] = writer._add_object(content)
    if encrypted:
        writer.encrypt('locked')
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_multipage_pdf_preserves_order_and_identifies_unreadable_middle_page():
    texts = ['First page: installed cable on the first floor today.', '',
             'Third page: installed pipe on the third floor today.']
    parts, warnings, kind = reports._parse_bytes_local(pdf_bytes(texts), 'diary.pdf', 'application/pdf')
    assert parts == [('page:1', texts[0]), ('page:2', ''), ('page:3', texts[2])]
    assert warnings == ['NEEDS_TRANSCRIPTION:page:2'] and kind == 'pdf'


def test_encrypted_and_oversized_page_count_pdfs_rejected():
    with pytest.raises(ValueError, match='encrypted PDF'):
        reports._parse_bytes_local(pdf_bytes(['locked'], encrypted=True), 'locked.pdf', 'application/pdf')
    with pytest.raises(ValueError, match='20 page'):
        reports._parse_bytes_local(pdf_bytes([''] * 21), 'many.pdf', 'application/pdf')


def test_spreadsheet_row_limit_cannot_return_partial_success(monkeypatch):
    monkeypatch.setattr(reports, 'MAX_XLSX_ROWS', 1)
    book = Workbook()
    book.active.append(['report_date', 'area', 'asset_tag', 'work_description', 'quantity', 'unit', 'quantity_kind'])
    for _ in range(2):
        book.active.append(['2026-10-20', 'A', 'P-1', 'Installed pipe', 2, 'm', 'delta'])
    raw = BytesIO()
    book.save(raw)
    book.close()
    with pytest.raises(ValueError, match='5,000 row'):
        reports._parse_bytes_local(raw.getvalue(), 'many.xlsx', 'application/octet-stream')


@pytest.mark.parametrize('prefix', ['=', '+', '-', '@', '\t=', '\r=', '\n=', ' \t@'])
def test_export_formula_prefixes_are_neutralized(prefix):
    from app.api.endpoints.progress import _safe_csv
    text = prefix + 'HYPERLINK("https://attacker.invalid", "click")'
    assert _safe_csv(text) == "'" + text


def test_tool_surface_has_no_file_network_or_execution_arguments():
    import inspect
    from app.agent.tools import AgentTools
    expected = {
        'load_job_context': {'job_id'},
        'extract_observations': {'fragment_ids', 'report_metadata'},
        'retrieve_candidates': {'observation_key', 'observation', 'schedule_version_id'},
        'select_candidate': {'observation_key', 'observation', 'candidates', 'fragment_ids'},
        'validate_proposal_drafts': {'state'},
        'persist_review_proposals': {'state'},
    }
    for name, arguments in expected.items():
        parameters = inspect.signature(getattr(AgentTools, name)).parameters
        assert set(parameters) - {'self'} == arguments
        assert all(p.kind not in (p.VAR_KEYWORD, p.VAR_POSITIONAL) for p in parameters.values())


def test_native_reader_treats_shell_syntax_in_identifiers_as_literal_argv(tmp_path, monkeypatch):
    import subprocess
    from app.ingest.native_schedule.runner import run_reader_module
    source = tmp_path / 'file; touch unauthorized.xer'
    source.write_text('fixture')
    seen = []
    def run(command, timeout):
        seen.append((command, timeout))
        return subprocess.CompletedProcess(command, 0, '{"task_count": 1}', '')
    monkeypatch.setattr('app.ingest.native_schedule.runner._capture_child', run)
    run_reader_module('app.ingest.native_schedule.primavera', source, '$(touch unauthorized)')
    command, timeout = seen[0]
    assert command[-1] == '$(touch unauthorized)'
    assert command[-3].endswith('/input.xer')
    assert ';' not in command[-3]
    assert timeout == 30
    assert not (tmp_path / 'unauthorized').exists()


@pytest.mark.parametrize('outside_kind', ['traversal', 'absolute', 'symlink'])
def test_original_download_rejects_paths_outside_upload_root(tmp_path, monkeypatch, outside_kind):
    from types import SimpleNamespace
    from uuid import uuid4
    from fastapi import HTTPException
    from app.api.endpoints import report_intake
    root = tmp_path / 'uploads'
    root.mkdir()
    outside = tmp_path / 'private.txt'
    outside.write_text('private')
    key = '../private.txt' if outside_kind == 'traversal' else str(outside)
    if outside_kind == 'symlink':
        (root / 'link').symlink_to(outside)
        key = 'link'
    project = SimpleNamespace(id=uuid4())
    record = SimpleNamespace(project_id=project.id, storage_key=key)
    monkeypatch.setattr(report_intake, 'get_settings', lambda: SimpleNamespace(upload_dir=str(root)))
    monkeypatch.setattr(report_intake, 'report_for', lambda *a: SimpleNamespace(file_id=uuid4()))
    database = SimpleNamespace(get=lambda *a: record)
    with pytest.raises(HTTPException) as error:
        report_intake.original(uuid4(), project, database, SimpleNamespace())
    assert error.value.status_code == 404


def test_html_original_download_is_attachment_with_nosniff(tmp_path, monkeypatch):
    from types import SimpleNamespace
    from uuid import uuid4
    from app.api.endpoints import report_intake
    project = SimpleNamespace(id=uuid4())
    (tmp_path / 'report.html').write_text('<script>alert(1)</script>')
    record = SimpleNamespace(project_id=project.id, storage_key='report.html', source_kind='text',
                             original_filename='report.html', mime_type='text/html')
    monkeypatch.setattr(report_intake, 'get_settings', lambda: SimpleNamespace(upload_dir=str(tmp_path)))
    monkeypatch.setattr(report_intake, 'report_for', lambda *a: SimpleNamespace(file_id=uuid4()))
    response = report_intake.original(uuid4(), project, SimpleNamespace(get=lambda *a: record), SimpleNamespace())
    assert response.media_type == 'application/octet-stream'
    assert response.headers['content-disposition'].startswith('attachment;')
    assert response.headers['x-content-type-options'] == 'nosniff'
